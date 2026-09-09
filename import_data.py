#!/usr/bin/env python3
"""將 Excel 學生帳號密碼資料匯入 SQLite 資料庫。"""

import sqlite3
import openpyxl
import os

EXCEL_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '115學生帳號密碼.xlsx')
DB_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'students.db')


def main():
    # 讀取 Excel（使用 data_only=True 取得計算後的值）
    print(f'正在讀取 {EXCEL_FILE} ...')
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    ws = wb['工作表1']

    # 若資料庫已存在則刪除重建
    if os.path.exists(DB_FILE):
        os.remove(DB_FILE)
        print('已刪除舊資料庫。')

    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()

    # 建立學生資料表
    cursor.execute('''
        CREATE TABLE students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            class_name TEXT NOT NULL,
            seat_number TEXT NOT NULL,
            name TEXT NOT NULL,
            iam_account TEXT,
            iam_password TEXT,
            cjes_account TEXT NOT NULL,
            cjes_password TEXT NOT NULL
        )
    ''')

    # 建立成績記錄表
    cursor.execute('''
        CREATE TABLE scores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id INTEGER NOT NULL,
            score INTEGER NOT NULL,
            rounds INTEGER NOT NULL,
            accuracy REAL NOT NULL,
            time_limit INTEGER NOT NULL,
            created_at DATETIME DEFAULT (datetime('now', 'localtime')),
            FOREIGN KEY (student_id) REFERENCES students(id)
        )
    ''')

    # 建立索引以加速查詢
    cursor.execute('CREATE INDEX idx_students_class ON students(class_name)')
    cursor.execute('CREATE INDEX idx_scores_student ON scores(student_id)')
    cursor.execute('CREATE INDEX idx_scores_time_limit ON scores(time_limit)')

    # 匯入資料
    count = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[0] is None:
            continue

        class_name = str(row[0]).strip()
        seat_number = str(row[1]).strip()
        name = str(row[2]).strip() if row[2] else ''
        iam_account = str(row[3]).strip() if row[3] else ''
        iam_password = str(row[4]).strip() if row[4] else ''
        cjes_account = str(row[5]).strip() if row[5] else ''
        cjes_password = str(row[6]).strip() if row[6] else ''

        if not name or not cjes_account:
            continue

        cursor.execute('''
            INSERT INTO students
                (class_name, seat_number, name, iam_account, iam_password, cjes_account, cjes_password)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', (class_name, seat_number, name, iam_account, iam_password, cjes_account, cjes_password))
        count += 1

    conn.commit()
    conn.close()
    print(f'✅ 成功匯入 {count} 筆學生資料到 {DB_FILE}')


if __name__ == '__main__':
    main()
